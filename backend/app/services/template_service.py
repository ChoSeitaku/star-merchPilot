"""导入模板生成（运行时生成 xlsx，不在仓库中存放二进制文件）。

每个模板包含：
- 「数据」工作表：表头 + 示例行 + 字段说明批注页
- 「填写说明」工作表：字段含义、必填项与格式要求
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F6FEB")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
NOTE_FONT = Font(color="B91C1C", size=10)


class Column:
    def __init__(
        self,
        key: str,
        title: str,
        *,
        required: bool = False,
        kind: str = "str",
        note: str = "",
        width: int = 18,
        sample: Any = "",
    ) -> None:
        self.key = key
        self.title = title
        self.required = required
        self.kind = kind
        self.note = note
        self.width = width
        self.sample = sample


# ------------------------------------------------------------------ 模板定义
PRODUCT_COLUMNS = [
    Column("store_name", "店铺名称", required=True,
           note="必填，需与库中已有店铺名称完全一致；也可用 store_id 参数指定默认店铺",
           width=22, sample="星航优选旗舰店"),
    Column("name", "商品名称", required=True, note="必填，2-255 字", width=30, sample="轻量通勤双肩包"),
    Column("platform", "平台", required=True, kind="enum",
           note="必填，可选：taobao/tmall/jd/pdd/douyin/xiaohongshu/other", width=14, sample="taobao"),
    Column("category", "类目", required=True, note="必填，例如：箱包/女包", width=16, sample="箱包"),
    Column("price", "售价", required=True, kind="number", note="必填，≥0 的数字", width=12, sample=199.0),
    Column("cost", "成本", kind="number", note="选填，≥0 的数字", width=12, sample=88.0),
    Column("target_audience", "目标人群", width=24, sample="22-35 岁通勤女性"),
    Column("selling_points", "核心卖点", width=32, note="多条用英文逗号分隔", sample="轻至480g,防泼水,可放14寸笔记本"),
    Column("product_url", "商品链接", width=32, sample="https://item.taobao.com/item.htm?id=10001"),
    Column("status", "商品状态", kind="enum", note="选填，draft/active/inactive，默认 draft", width=12, sample="draft"),
]

SKU_COLUMNS = [
    Column("product_name", "商品名称", required=True, note="必填，需与库中已有商品名称完全一致", width=30, sample="轻量通勤双肩包"),
    Column("sku_code", "SKU编码", required=True, kind="text",
           note="必填且同一商品下唯一；请把该列设置为「文本」格式，避免 00123 被识别成 123",
           width=18, sample="BAG-BLK-001"),
    Column("sku_name", "SKU名称", width=24, sample="黑色标准款"),
    Column("spec", "规格", width=24, note="JSON 或 颜色:黑色;尺寸:标准", sample="颜色:黑色;尺寸:标准"),
    Column("price", "售价", kind="number", note="选填，留空则继承商品售价", width=12, sample=199.0),
    Column("cost", "成本", kind="number", width=12, sample=88.0),
    Column("stock_qty", "库存数量", kind="int", note="整数，≥0", width=12, sample=120),
    Column("warning_threshold", "预警阈值", kind="int", note="整数，≥0；低于该值会触发低库存提示", width=12, sample=15),
    Column("location_text", "库位", width=16, sample="A区-03-12"),
]

PERFORMANCE_COLUMNS = [
    Column("period_start", "周期开始", required=True, kind="date", note="必填，格式 YYYY-MM-DD", width=14, sample="2026-09-01"),
    Column("period_end", "周期结束", required=True, kind="date", note="必填，格式 YYYY-MM-DD，不早于开始日期", width=14, sample="2026-09-07"),
    Column("impressions", "曝光量", kind="int", note="整数，≥0", width=12, sample=120000),
    Column("clicks", "点击量", kind="int", note="整数，≥0，不能大于曝光量", width=12, sample=3600),
    Column("conversions", "成交量", kind="int", note="整数，≥0，不能大于点击量", width=12, sample=72),
    Column("spend", "花费金额", kind="number", note="≥0 的数字", width=12, sample=1800.0),
    Column("revenue", "成交金额", kind="number", note="≥0 的数字", width=12, sample=6480.0),
    Column("ctr", "点击率(%)", kind="number", note="选填；留空则自动按 点击/曝光 计算", width=12, sample=None),
    Column("conversion_rate", "转化率(%)", kind="number", note="选填；留空则自动按 成交/点击 计算", width=12, sample=None),
    Column("roi", "ROI", kind="number", note="选填；留空则自动按 成交金额/花费 计算", width=10, sample=None),
    Column("notes", "备注", width=26, sample="国庆预热期，主图方向一上线"),
]

MAPPING_COLUMNS = [
    Column("product_name", "商品名称", required=True, note="必填，需与库中已有商品名称完全一致", width=30, sample="轻量通勤双肩包"),
    Column("platform", "平台", required=True, kind="enum",
           note="必填，taobao/tmall/jd/pdd/douyin/xiaohongshu/other", width=14, sample="taobao"),
    Column("platform_product_id", "平台商品ID", required=True, note="必填", width=22, sample="678901234567"),
    Column("platform_sku_id", "平台SKU编码", width=22, sample="SKU-998877"),
    Column("mapping_status", "映射状态", kind="enum", note="mapped/unmapped/conflict，默认 mapped", width=14, sample="mapped"),
]

TEMPLATES: dict[str, dict[str, Any]] = {
    "products": {
        "filename": "商品导入模板.xlsx",
        "title": "商品导入模板",
        "columns": PRODUCT_COLUMNS,
        "description": "用于批量导入商品主档信息。导出后请勿修改表头行。",
    },
    "sku-inventory": {
        "filename": "SKU库存导入模板.xlsx",
        "title": "SKU 与库存导入模板",
        "columns": SKU_COLUMNS,
        "description": "用于为已有商品批量导入 SKU 与期初库存。商品需先存在。",
    },
    "performance-records": {
        "filename": "经营数据导入模板.xlsx",
        "title": "经营数据导入模板",
        "columns": PERFORMANCE_COLUMNS,
        "description": "用于按周期导入投放经营数据，导入后可用于生成复盘报告。",
    },
    "platform-mappings": {
        "filename": "平台字段映射导入模板.xlsx",
        "title": "平台字段映射导入模板",
        "columns": MAPPING_COLUMNS,
        "description": "用于把系统商品与各平台商品 ID 建立映射关系。",
    },
}


# 别名：导入接口用 `performance`，模板键用 `performance-records`，两者都接受
TEMPLATE_ALIASES = {
    "performance": "performance-records",
    "sku_inventory": "sku-inventory",
    "sku": "sku-inventory",
    "platform_mappings": "platform-mappings",
    "mappings": "platform-mappings",
}


def resolve_type(import_type: str) -> str:
    """把别名归一化为模板键。"""
    return TEMPLATE_ALIASES.get(import_type, import_type)


# ------------------------------------------------------------------ 生成文件
def build_template(import_type: str) -> tuple[BytesIO, str]:
    """生成模板文件，返回 (字节流, 下载文件名)。"""
    spec = TEMPLATES.get(resolve_type(import_type))
    if spec is None:
        raise KeyError(f"未知的模板类型：{import_type}")

    columns: list[Column] = spec["columns"]
    wb = Workbook()

    ws = wb.active
    ws.title = "数据"
    for idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx, value=col.title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = col.width

        title_cell = ws.cell(row=1, column=idx)
        comment_bits = []
        if col.required:
            comment_bits.append("【必填】")
        if col.note:
            comment_bits.append(col.note)
        if comment_bits:
            title_cell.comment = _make_comment(" ".join(comment_bits))

    # 示例行
    for idx, col in enumerate(columns, start=1):
        value = col.sample
        cell = ws.cell(row=2, column=idx, value=value)
        if col.kind == "text" and value not in ("", None):
            cell.number_format = "@"
        cell.font = Font(color="6B7280", italic=True)

    # 必填列标记
    for idx, col in enumerate(columns, start=1):
        if col.required:
            ws.cell(row=1, column=idx).value = f"{col.title} *"

    ws.freeze_panes = "A2"

    # 填写说明
    guide = wb.create_sheet("填写说明")
    guide.column_dimensions["A"].width = 16
    guide.column_dimensions["B"].width = 12
    guide.column_dimensions["C"].width = 56
    guide.cell(row=1, column=1, value=spec["title"]).font = Font(bold=True, size=13)
    guide.cell(row=2, column=1, value=spec["description"]).font = NOTE_FONT
    guide.cell(row=4, column=1, value="字段").font = Font(bold=True)
    guide.cell(row=4, column=2, value="是否必填").font = Font(bold=True)
    guide.cell(row=4, column=3, value="说明").font = Font(bold=True)
    for i, col in enumerate(columns, start=5):
        guide.cell(row=i, column=1, value=col.title)
        guide.cell(row=i, column=2, value="必填" if col.required else "选填")
        guide.cell(row=i, column=3, value=col.note or "—")

    row = len(columns) + 6
    guide.cell(row=row, column=1, value="导入规则").font = Font(bold=True)
    rules = [
        "1. 请勿修改或删除表头行；标有 * 的列为必填。",
        "2. 系统会先做字段校验并返回预览与错误行提示，确认后才写入数据库。",
        "3. 导入允许部分成功：校验通过的行会正常写入，出错的行会被跳过并给出原因。",
        "4. SKU 编码等可能以 0 开头的字段，请把单元格格式设置为「文本」。",
        "5. 日期字段请使用 YYYY-MM-DD 格式，例如 2026-09-01。",
    ]
    for offset, rule in enumerate(rules, start=1):
        guide.cell(row=row + offset, column=1, value=rule)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer, spec["filename"]


def _make_comment(text: str):
    from openpyxl.comments import Comment

    return Comment(text, "电商运营助手")


def template_columns(import_type: str) -> list[Column]:
    return TEMPLATES[resolve_type(import_type)]["columns"]
