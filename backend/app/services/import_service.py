"""导入服务：模板校验 → 预览 → 确认导入（允许部分成功）。

两段式流程：
1. ``preview_only=True``：解析上传文件，逐行校验，把文件暂存到 ``storage/tmp``，
   返回总行数 / 有效行数 / 错误行明细 / 样例行。**不写库**。
2. ``preview_only=False``：用 ``import_id`` 取回暂存文件重新解析，
   逐行在 savepoint 内提交，单行失败不影响其他行，返回成功与失败统计。

日期与数字的宽松解析集中在这里，保证前端拿到的错误行信息足够定位问题。
"""

from __future__ import annotations

import json
import uuid
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Optional

from fastapi import HTTPException, UploadFile, status
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from ..config import settings
from ..database import utcnow
from ..models import (
    Store,
    PLATFORMS,
    PRODUCT_STATUSES,

    InventoryItem,
    InventoryMovement,
    PlatformProductMapping,
    Product,
    ProductSku,
    PerformanceRecord,
)
from . import template_service
from .template_service import Column

MAX_ROWS = 2000
PREVIEW_SAMPLE_LIMIT = 8


# --------------------------------------------------------------- 值解析工具
class RowError(Exception):
    def __init__(self, messages: list[str]) -> None:
        super().__init__("; ".join(messages))
        self.messages = messages


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        # Excel 常把整数读成 100.0，这里还原成 100，避免 SKU 编码出现 .0
        return str(int(value))
    return str(value).strip()


def _parse_number(value: Any, title: str, errors: list[str], *, required: bool) -> Optional[float]:
    text = _cell_str(value)
    if not text:
        if required:
            errors.append(f"{title} 为必填项")
        return None
    try:
        number = float(text)
    except ValueError:
        errors.append(f"{title} 应为数字，实际为「{text}」")
        return None
    if number < 0:
        errors.append(f"{title} 不能为负数，实际为 {number}")
        return None
    return number


def _parse_int(value: Any, title: str, errors: list[str], *, required: bool) -> Optional[int]:
    number = _parse_number(value, title, errors, required=required)
    if number is None:
        return None
    if abs(number - round(number)) > 1e-9:
        errors.append(f"{title} 应为整数，实际为 {number}")
        return None
    return int(round(number))


def _parse_date(value: Any, title: str, errors: list[str], *, required: bool) -> Optional[date]:
    text = _cell_str(value)
    if not text:
        if required:
            errors.append(f"{title} 为必填项")
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    errors.append(f"{title} 格式应为 YYYY-MM-DD，实际为「{text}」")
    return None


# --------------------------------------------------------------- 表单读取
def _read_rows(content: bytes, import_type: str) -> tuple[list[Column], list[dict[str, Any]]]:
    """读取 xlsx，返回 (列定义, 行数据列表)。行数据以列 key 为键。"""
    try:
        wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"无法解析文件，请确认是有效的 xlsx 文件：{exc}",
        ) from exc

    ws = wb["数据"] if "数据" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="文件为空") from None

    columns = template_service.template_columns(import_type)
    header_titles = [_cell_str(h).replace(" *", "").strip() for h in header]
    title_to_index = {title: idx for idx, title in enumerate(header_titles) if title}

    # 表头允许乱序，但必须能匹配到模板字段
    missing = [c.title for c in columns if c.required and c.title not in title_to_index]
    if missing:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"缺少必填列：{'、'.join(missing)}；请使用系统提供的导入模板",
        )

    parsed: list[dict[str, Any]] = []
    for row_no, raw in enumerate(rows_iter, start=2):
        if raw is None or all(v is None or _cell_str(v) == "" for v in raw):
            continue  # 跳过空行
        item: dict[str, Any] = {"__row_no__": row_no}
        for col in columns:
            idx = title_to_index.get(col.title)
            item[col.key] = raw[idx] if idx is not None and idx < len(raw) else None
        parsed.append(item)
        if len(parsed) > MAX_ROWS:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"单次导入最多 {MAX_ROWS} 行，请拆分文件后重试",
            )
    return columns, parsed


# --------------------------------------------------------------- 各类校验
def _validate_product_row(
    db: Session, row: dict[str, Any], seen_names: set[str], store_lookup: dict[str, "Store"]
) -> dict[str, Any]:
    errors: list[str] = []
    name = _cell_str(row.get("name"))
    if not name:
        errors.append("商品名称为必填项")
    elif len(name) < 2:
        errors.append("商品名称至少 2 个字符")
    elif name in seen_names:
        errors.append(f"商品名称「{name}」在本次文件中重复")
    seen_names.add(name)

    store_name = _cell_str(row.get("store_name"))
    store = store_lookup.get(store_name)
    if not store_name:
        errors.append("店铺名称为必填项")
    elif store is None:
        errors.append(f"未找到店铺「{store_name}」，请先创建该店铺")

    platform = _cell_str(row.get("platform")) or "other"
    if platform not in PLATFORMS:
        errors.append(f"平台「{platform}」不在允许范围（{'/'.join(PLATFORMS)}）")

    category = _cell_str(row.get("category"))
    if not category:
        errors.append("类目为必填项")

    price = _parse_number(row.get("price"), "售价", errors, required=True)
    cost = _parse_number(row.get("cost"), "成本", errors, required=False)
    if price is not None and cost is not None and cost > price:
        errors.append(f"成本（{cost}）高于售价（{price}），请确认是否填错")

    product_status = _cell_str(row.get("status")) or "draft"
    if product_status not in PRODUCT_STATUSES:
        errors.append(f"商品状态「{product_status}」不在允许范围（{'/'.join(PRODUCT_STATUSES)}）")

    if errors:
        raise RowError(errors)

    return {
        "store": store,
        "name": name,
        "platform": platform,
        "category": category,
        "price": price or 0.0,
        "cost": cost or 0.0,
        "target_audience": _cell_str(row.get("target_audience")),
        "selling_points": _cell_str(row.get("selling_points")),
        "product_url": _cell_str(row.get("product_url")),
        "status": product_status,
    }


def _validate_sku_row(
    db: Session,
    row: dict[str, Any],
    product_lookup: dict[str, Product],
    seen_codes: dict[str, set[str]],
) -> dict[str, Any]:
    errors: list[str] = []
    product_name = _cell_str(row.get("product_name"))
    product = product_lookup.get(product_name)
    if not product_name:
        errors.append("商品名称为必填项")
    elif product is None:
        errors.append(f"未找到商品「{product_name}」，请先导入或创建该商品")

    sku_code = _cell_str(row.get("sku_code"))
    if not sku_code:
        errors.append("SKU 编码为必填项")
    elif product is not None:
        if sku_code in seen_codes.setdefault(product_name, set()):
            errors.append(f"SKU 编码「{sku_code}」在本次文件中重复")
        seen_codes[product_name].add(sku_code)
        exists = db.scalar(
            select(ProductSku.id).where(
                ProductSku.product_id == product.id, ProductSku.sku_code == sku_code
            )
        )
        if exists:
            errors.append(f"SKU 编码「{sku_code}」在商品「{product_name}」下已存在")

    price = _parse_number(row.get("price"), "售价", errors, required=False)
    cost = _parse_number(row.get("cost"), "成本", errors, required=False)
    stock = _parse_int(row.get("stock_qty"), "库存数量", errors, required=False)
    threshold = _parse_int(row.get("warning_threshold"), "预警阈值", errors, required=False)

    spec_raw = _cell_str(row.get("spec"))
    spec: dict[str, Any] = {}
    if spec_raw:
        try:
            parsed_spec = json.loads(spec_raw)
            if isinstance(parsed_spec, dict):
                spec = parsed_spec
            else:
                raise ValueError
        except (ValueError, json.JSONDecodeError):
            # 兼容 `颜色:黑色;尺寸:标准` 这种写法
            for pair in spec_raw.replace("；", ";").split(";"):
                if ":" in pair or "：" in pair:
                    k, _, v = pair.replace("：", ":").partition(":")
                    if k.strip():
                        spec[k.strip()] = v.strip()
            if not spec:
                errors.append(f"规格「{spec_raw}」无法解析，请使用 JSON 或 颜色:黑色;尺寸:标准 格式")

    if errors:
        raise RowError(errors)

    return {
        "product": product,
        "sku_code": sku_code,
        "sku_name": _cell_str(row.get("sku_name")) or sku_code,
        "spec_json": spec,
        "price": price if price is not None else (product.price if product else 0.0),
        "cost": cost if cost is not None else (product.cost if product else 0.0),
        "stock_qty": stock or 0,
        "warning_threshold": threshold if threshold is not None else 10,
        "location_text": _cell_str(row.get("location_text")),
    }


def _validate_performance_row(db: Session, row: dict[str, Any], product: Product) -> dict[str, Any]:
    errors: list[str] = []
    period_start = _parse_date(row.get("period_start"), "周期开始", errors, required=True)
    period_end = _parse_date(row.get("period_end"), "周期结束", errors, required=True)

    impressions = _parse_int(row.get("impressions"), "曝光量", errors, required=False) or 0
    clicks = _parse_int(row.get("clicks"), "点击量", errors, required=False) or 0
    conversions = _parse_int(row.get("conversions"), "成交量", errors, required=False) or 0
    spend = _parse_number(row.get("spend"), "花费金额", errors, required=False) or 0.0
    revenue = _parse_number(row.get("revenue"), "成交金额", errors, required=False) or 0.0

    if period_start and period_end and period_end < period_start:
        errors.append(f"周期结束（{period_end}）早于周期开始（{period_start}）")
    if impressions and clicks > impressions:
        errors.append(f"点击量（{clicks}）不能大于曝光量（{impressions}）")
    if clicks and conversions > clicks:
        errors.append(f"成交量（{conversions}）不能大于点击量（{clicks}）")
    if clicks and not impressions:
        errors.append("填写了点击量但曝光量为空，请补全曝光量")

    ctr = _parse_number(row.get("ctr"), "点击率", errors, required=False)
    cvr = _parse_number(row.get("conversion_rate"), "转化率", errors, required=False)
    roi = _parse_number(row.get("roi"), "ROI", errors, required=False)

    if errors:
        raise RowError(errors)

    return {
        "product_id": product.id,
        "period_start": period_start,
        "period_end": period_end,
        "impressions": impressions,
        "clicks": clicks,
        "conversions": conversions,
        "spend": spend,
        "revenue": revenue,
        # 未提供则按通用口径自动计算
        "ctr": ctr if ctr is not None else (round(clicks / impressions * 100, 4) if impressions else 0.0),
        "conversion_rate": (
            cvr if cvr is not None else (round(conversions / clicks * 100, 4) if clicks else 0.0)
        ),
        "roi": roi if roi is not None else (round(revenue / spend, 4) if spend else 0.0),
        "notes": _cell_str(row.get("notes")),
    }


def _validate_mapping_row(
    db: Session, row: dict[str, Any], product_lookup: dict[str, Product]
) -> dict[str, Any]:
    errors: list[str] = []
    product_name = _cell_str(row.get("product_name"))
    product = product_lookup.get(product_name)
    if not product_name:
        errors.append("商品名称为必填项")
    elif product is None:
        errors.append(f"未找到商品「{product_name}」")

    platform = _cell_str(row.get("platform"))
    if platform not in PLATFORMS:
        errors.append(f"平台「{platform}」不在允许范围（{'/'.join(PLATFORMS)}）")

    platform_product_id = _cell_str(row.get("platform_product_id"))
    if not platform_product_id:
        errors.append("平台商品ID 为必填项")

    mapping_status = _cell_str(row.get("mapping_status")) or "mapped"
    if mapping_status not in ("mapped", "unmapped", "conflict"):
        errors.append(f"映射状态「{mapping_status}」不在允许范围（mapped/unmapped/conflict）")

    if errors:
        raise RowError(errors)

    return {
        "product": product,
        "platform": platform,
        "platform_product_id": platform_product_id,
        "platform_sku_id": _cell_str(row.get("platform_sku_id")),
        "mapping_status": mapping_status,
    }


# --------------------------------------------------------------- 导入主流程
IMPORT_LABELS = {
    "products": "商品导入",
    "sku-inventory": "SKU 库存导入",
    "performance": "经营数据导入",
    "performance-records": "经营数据导入",
    "platform-mappings": "平台字段映射导入",
}


def _tmp_path(import_id: str) -> Path:
    return settings.tmp_dir / f"{import_id}.xlsx"


def _build_preview_payload(
    import_type: str, total: int, valid: int, errors: list[dict], sample: list[dict]
) -> dict:
    return {
        "import_id": "",
        "import_type": import_type,
        "import_label": IMPORT_LABELS.get(import_type, import_type),
        "total_rows": total,
        "valid_rows": valid,
        "invalid_rows": errors,
        "error_rows": len(errors),
        "sample_rows": sample,
        "can_confirm": valid > 0,
    }


async def handle_import(
    db: Session,
    *,
    import_type: str,
    file: UploadFile,
    preview_only: bool,
    import_id: str = "",
    product_id: Optional[int] = None,
) -> dict:
    """统一的导入入口（预览与确认都由这里分发）。"""
    import_type = template_service.resolve_type(import_type)
    if import_type not in template_service.TEMPLATES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail=f"不支持的导入类型：{import_type}"
        )

    if preview_only:
        content = await file.read()
        if not content:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="上传文件为空")
        return _do_preview(db, import_type, content, product_id)

    if not import_id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="确认导入需要提供预览阶段返回的 import_id",
        )
    path = _tmp_path(import_id)
    if not path.exists():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="预览文件已过期或不存在，请重新上传并预览",
        )
    content = path.read_bytes()
    result = _do_commit(db, import_type, content, import_id, product_id)
    path.unlink(missing_ok=True)
    return result


def _do_preview(
    db: Session, import_type: str, content: bytes, product_id: Optional[int]
) -> dict:
    _, rows = _read_rows(content, import_type)
    validator = _validators(db, import_type, product_id)

    errors: list[dict] = []
    samples: list[dict] = []
    valid = 0
    for row in rows:
        row_no = row["__row_no__"]
        try:
            parsed = validator(row)
        except RowError as exc:
            errors.append({"row_no": row_no, "errors": exc.messages})
            continue
        valid += 1
        if len(samples) < PREVIEW_SAMPLE_LIMIT:
            samples.append(_sample_of(parsed))

    import_id = uuid.uuid4().hex
    _tmp_path(import_id).write_bytes(content)
    # 清理超过 2 小时的陈旧暂存文件
    _cleanup_tmp()

    payload = _build_preview_payload(import_type, len(rows), valid, errors, samples)
    payload["import_id"] = import_id
    payload["hint"] = (
        f"校验完成：共 {len(rows)} 行，{valid} 行可导入，{len(errors)} 行存在问题。"
        "确认后将只导入校验通过的行（允许部分成功）。"
    )
    return payload


def _do_commit(
    db: Session,
    import_type: str,
    content: bytes,
    import_id: str,
    product_id: Optional[int],
) -> dict:
    _, rows = _read_rows(content, import_type)
    validator = _validators(db, import_type, product_id)

    errors: list[dict] = []
    created_ids: list[int] = []
    success = 0

    for row in rows:
        row_no = row["__row_no__"]
        try:
            parsed = validator(row)
        except RowError as exc:
            errors.append({"row_no": row_no, "errors": exc.messages})
            continue

        # 每行一个 savepoint：单行失败不会影响已成功的行
        savepoint = db.begin_nested()
        try:
            created = _persist_row(db, import_type, parsed)
            savepoint.commit()
            success += 1
            if created is not None and getattr(created, "id", None):
                created_ids.append(created.id)
        except (IntegrityError, ValueError, KeyError) as exc:
            savepoint.rollback()
            errors.append({"row_no": row_no, "errors": [f"写入失败：{exc}"]})
        except Exception as exc:  # noqa: BLE001
            savepoint.rollback()
            errors.append({"row_no": row_no, "errors": [f"写入异常：{exc}"]})

    db.commit()
    return {
        "import_id": import_id,
        "import_type": import_type,
        "import_label": IMPORT_LABELS.get(import_type, import_type),
        "total_rows": len(rows),
        "success_rows": success,
        "failed_rows": len(errors),
        "errors": errors,
        "created_ids": created_ids,
        "message": (
            f"导入完成：成功 {success} 行，失败 {len(errors)} 行"
            + ("（部分成功）" if success and errors else "")
        ),
    }


# --------------------------------------------------------------- 校验器分发
def _validators(
    db: Session, import_type: str, product_id: Optional[int]
) -> Callable[[dict[str, Any]], dict[str, Any]]:
    # 归一化别名（例如 performance → performance-records）
    import_type = template_service.resolve_type(import_type)
    if import_type == "products":
        stores = list(db.scalars(select(Store)))
        store_lookup = {s.store_name: s for s in stores}
        seen_names: set[str] = set()
        return lambda row: _validate_product_row(db, row, seen_names, store_lookup)

    if import_type == "sku-inventory":
        products = list(db.scalars(select(Product)))
        lookup = {p.name: p for p in products}
        seen_codes: dict[str, set[str]] = {}
        return lambda row: _validate_sku_row(db, row, lookup, seen_codes)

    if import_type == "performance-records":
        if product_id is None:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST, detail="经营数据导入需要指定商品"
            )
        product = db.get(Product, product_id)
        if product is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="商品不存在")
        return lambda row: _validate_performance_row(db, row, product)

    if import_type == "platform-mappings":
        products = list(db.scalars(select(Product)))
        lookup = {p.name: p for p in products}
        return lambda row: _validate_mapping_row(db, row, lookup)

    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST, detail=f"不支持的导入类型：{import_type}"
    )


# --------------------------------------------------------------- 落库
def _persist_row(db: Session, import_type: str, parsed: dict[str, Any]):
    if import_type == "products":
        store: Store = parsed.pop("store")
        product = Product(store_id=store.id, **parsed, images_json=[])
        db.add(product)
        db.flush()
        return product

    if import_type == "sku-inventory":
        product: Product = parsed.pop("product")
        stock_qty = parsed.pop("stock_qty")
        warning_threshold = parsed.pop("warning_threshold")
        location_text = parsed.pop("location_text")

        sku = ProductSku(product_id=product.id, **parsed)
        db.add(sku)
        db.flush()
        db.add(
            InventoryItem(
                sku_id=sku.id,
                stock_qty=stock_qty,
                locked_qty=0,
                warning_threshold=warning_threshold,
                location_text=location_text,
            )
        )
        if stock_qty:
            db.add(
                InventoryMovement(
                    sku_id=sku.id,
                    movement_type="init",
                    change_qty=stock_qty,
                    before_qty=0,
                    after_qty=stock_qty,
                    reason_text="SKU 库存导入期初库存",
                    reference_type="import",
                    reference_id=None,
                )
            )
        return sku

    if import_type == "performance-records":
        record = PerformanceRecord(**parsed)
        db.add(record)
        db.flush()
        return record

    if import_type == "platform-mappings":
        product: Product = parsed.pop("product")
        mapping = PlatformProductMapping(product_id=product.id, raw_payload_json={"source": "import"}, **parsed)
        db.add(mapping)
        db.flush()
        return mapping

    raise ValueError(f"不支持的导入类型：{import_type}")


def _sample_of(parsed: dict[str, Any]) -> dict[str, Any]:
    """把已校验的行转成可 JSON 化的预览样例（ORM 对象替换为名称）。"""
    out: dict[str, Any] = {}
    for key, value in parsed.items():
        if isinstance(value, (Product, Store)):
            out[f"{key}_name"] = value.name if isinstance(value, Product) else value.store_name
        elif isinstance(value, (date, datetime)):
            out[key] = value.isoformat()
        else:
            out[key] = value
    return out


def _cleanup_tmp(max_age_hours: int = 2) -> None:
    """清理过期的暂存文件，避免磁盘堆积。"""
    cutoff = utcnow().timestamp() - max_age_hours * 3600
    for path in settings.tmp_dir.glob("*.xlsx"):
        try:
            if path.stat().st_mtime < cutoff:
                path.unlink(missing_ok=True)
        except OSError:
            continue
