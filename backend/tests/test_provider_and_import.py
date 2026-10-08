"""内容生成契约与导入校验测试。"""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.providers import get_ai_provider
from app.services import template_service


# =============================================================== 内容生成契约
DIAGNOSIS_FIELDS = [
    "positioning",
    "price_band",
    "audience_insights",
    "pain_points",
    "selling_point_analysis",
    "risks",
    "recommendations",
    "summary_text",
]

MAIN_IMAGE_DIRECTION_FIELDS = [
    "direction_title",
    "visual_structure",
    "core_copy",
    "highlight_selling_point",
    "rationale",
]

VIDEO_SCRIPT_FIELDS = ["script_title", "hook", "storyboard", "voiceover", "conversion_cta"]


@pytest.fixture()
def payload() -> dict:
    return {
        "product": {
            "id": 42,
            "name": "轻量通勤双肩包",
            "category": "箱包",
            "price": 199.0,
            "cost": 86.0,
            "target_audience": "通勤上班族",
            "selling_points": "轻量,防泼水,大容量",
        },
        "skus": [{"sku_code": "A", "stock_qty": 5, "warning_threshold": 10}],
        "competitors": [
            {"id": 1, "name": "竞品A", "price": 189.0},
            {"id": 2, "name": "竞品B", "price": 219.0},
        ],
        "performance": [
            {"id": 1, "impressions": 100000, "clicks": 3000, "ctr": 3.0,
             "conversions": 60, "conversion_rate": 2.0, "spend": 1500, "revenue": 5400, "roi": 3.6}
        ],
        "low_stock_count": 1,
        "prior_actions": ["重跑商品诊断：更新价格带分析"],
    }


def test_diagnosis_has_all_required_fields(payload):
    output = get_ai_provider().generate_diagnosis(payload)
    for field in DIAGNOSIS_FIELDS:
        assert output.get(field), f"诊断缺少字段 {field}"


def test_diagnosis_is_data_driven(payload):
    """输出应引用真实商品数据，而不是泛泛而谈。"""
    output = get_ai_provider().generate_diagnosis(payload)
    text = " ".join(str(v) for v in output.values())
    assert "轻量通勤双肩包" in text
    assert "199" in text
    # 竞品价格区间应被引用
    assert "189" in text and "219" in text


def test_diagnosis_includes_review_actions(payload):
    output = get_ai_provider().generate_diagnosis(payload)
    assert "上一轮复盘遗留动作" in output["recommendations"]


def test_diagnosis_is_deterministic(payload):
    """同一商品重复生成结果一致（便于重试复现与测试）。"""
    a = get_ai_provider().generate_diagnosis(payload)
    b = get_ai_provider().generate_diagnosis(payload)
    assert a == b


def test_main_image_plans_at_least_three_directions(payload):
    output = get_ai_provider().generate_main_image_plans(payload)
    directions = output["directions"]
    assert len(directions) >= 3
    for direction in directions:
        for field in MAIN_IMAGE_DIRECTION_FIELDS:
            assert direction.get(field), f"主图方向缺少字段 {field}"


def test_video_scripts_at_least_three(payload):
    output = get_ai_provider().generate_video_scripts(payload)
    scripts = output["scripts"]
    assert len(scripts) >= 3
    for script in scripts:
        for field in VIDEO_SCRIPT_FIELDS:
            assert script.get(field), f"视频脚本缺少字段 {field}"


def test_ad_recommendation_fields(payload):
    payload.update({"assets": [], "links": []})
    output = get_ai_provider().generate_ad_recommendation(payload)
    for field in [
        "summary_text",
        "objective_text",
        "audience_segments",
        "budget_plan",
        "creative_tests",
        "bid_strategy",
        "risk_controls",
        "next_steps",
    ]:
        assert output.get(field), f"投放建议缺少字段 {field}"
    # 合规：必须显式声明不自动执行
    joined = " ".join(output["risk_controls"])
    assert "不自动" in joined


def test_review_report_fields(payload):
    payload.update({"period_start": "2026-09-01", "period_end": "2026-09-30",
                    "plans": [], "diagnosis": {"positioning": "中高价位"}})
    output = get_ai_provider().generate_review_report(payload)
    assert output["summary_text"]
    assert output["insights"]
    assert output["next_actions"]
    for action in output["next_actions"]:
        assert action.get("action_title")
        assert action.get("action_detail")


def test_inventory_advice_flags_low_stock():
    output = get_ai_provider().generate_inventory_advice(
        {
            "store_name": "测试店铺",
            "items": [
                {"sku_code": "A", "sku_name": "低库存款", "stock_qty": 3, "warning_threshold": 10},
                {"sku_code": "B", "sku_name": "正常款", "stock_qty": 100, "warning_threshold": 10},
            ],
        }
    )
    assert "1 个低于预警阈值" in output["summary_text"]
    # 低库存的 SKU 应排在前面
    assert output["items"][0]["sku_code"] == "A"
    assert output["items"][0]["suggested_action"] == "建议尽快补货"


def test_promotion_link_suggestions():
    suggestions = get_ai_provider().generate_promotion_link_suggestions(
        {"product": {"id": 1, "name": "测试商品", "category": "箱包",
                     "product_url": "https://example.com/i/1"}}
    )
    assert len(suggestions) >= 3
    for item in suggestions:
        assert item["link_name"]
        assert item["target_url"]
        assert item["utm"]


# =============================================================== 模板与导入
def test_all_templates_build_and_reopen():
    """四个模板都能生成，且能被 openpyxl 重新打开。"""
    for import_type in ["products", "sku-inventory", "performance-records", "platform-mappings"]:
        buffer, filename = template_service.build_template(import_type)
        assert filename.endswith(".xlsx")
        wb = load_workbook(BytesIO(buffer.getvalue()))
        assert "数据" in wb.sheetnames
        assert "填写说明" in wb.sheetnames
        header = [c.value for c in wb["数据"][1]]
        assert all(header), f"{import_type} 表头不应有空列"


def test_template_alias_resolves():
    assert template_service.resolve_type("performance") == "performance-records"
    assert template_service.resolve_type("sku") == "sku-inventory"
    assert template_service.resolve_type("products") == "products"


def test_required_columns_marked():
    for import_type in ["products", "sku-inventory", "performance-records"]:
        columns = template_service.template_columns(import_type)
        assert any(c.required for c in columns), f"{import_type} 应至少有一个必填列"


# --------------------------------------------------------------- 导入接口
def _xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    for row in rows:
        ws.append(row)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


PRODUCT_HEADER = [
    "店铺名称 *", "商品名称 *", "平台 *", "类目 *", "售价 *", "成本",
    "目标人群", "核心卖点", "商品链接", "商品状态",
]


def test_product_import_preview_reports_errors(client, operator_headers, sample_store):
    rows = [
        PRODUCT_HEADER,
        ["测试店铺", "有效商品A", "taobao", "箱包", 159.0, 70.0, "", "", "", "draft"],
        ["不存在的店铺", "无效商品", "taobao", "箱包", 100.0, 50.0, "", "", "", "draft"],
        ["测试店铺", "平台非法", "不存在的平台", "箱包", 100.0, 50.0, "", "", "", "draft"],
        ["测试店铺", "成本高于售价", "taobao", "箱包", 100.0, 200.0, "", "", "", "draft"],
    ]
    response = client.post(
        "/api/v1/imports/products/preview",
        files={"file": ("p.xlsx", _xlsx(rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=operator_headers,
    )
    assert response.status_code == 200, response.text
    data = response.json()

    assert data["total_rows"] == 4
    assert data["valid_rows"] == 1
    assert data["error_rows"] == 3
    assert len(data["invalid_rows"]) == 3
    for row in data["invalid_rows"]:
        assert row["row_no"] >= 2
        assert row["errors"]

    messages = " ".join(e for row in data["invalid_rows"] for e in row["errors"])
    assert "未找到店铺" in messages
    assert "不在允许范围" in messages
    assert "高于售价" in messages


def test_product_import_partial_success(client, operator_headers, sample_store):
    """允许部分成功：有效行写入，错误行跳过。"""
    rows = [
        PRODUCT_HEADER,
        ["测试店铺", "导入成功商品A", "taobao", "箱包", 159.0, 70.0, "", "", "", "draft"],
        ["测试店铺", "导入成功商品B", "jd", "箱包", 259.0, 120.0, "", "", "", "active"],
        ["测试店铺", "X", "taobao", "箱包", 100.0, 50.0, "", "", "", "draft"],  # 名称过短
    ]
    preview = client.post(
        "/api/v1/imports/products/preview",
        files={"file": ("p.xlsx", _xlsx(rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=operator_headers,
    ).json()

    result = client.post(
        f"/api/v1/imports/products/confirm?import_id={preview['import_id']}",
        headers=operator_headers,
    )
    assert result.status_code == 200, result.text
    data = result.json()
    assert data["success_rows"] == 2
    assert data["failed_rows"] == 1
    assert len(data["created_ids"]) == 2
    assert "部分成功" in data["message"]


def test_import_confirm_requires_valid_import_id(client, operator_headers):
    response = client.post(
        "/api/v1/imports/products/confirm?import_id=deadbeef", headers=operator_headers
    )
    assert response.status_code == 404


def test_preview_does_not_persist(client, operator_headers, sample_store, db):
    from sqlalchemy import select

    from app.models import Product

    before = len(list(db.scalars(select(Product))))
    rows = [PRODUCT_HEADER, ["测试店铺", "不应落库的商品", "taobao", "箱包", 10.0, 5.0, "", "", "", "draft"]]
    client.post(
        "/api/v1/imports/products/preview",
        files={"file": ("p.xlsx", _xlsx(rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=operator_headers,
    )
    assert len(list(db.scalars(select(Product)))) == before


def test_performance_import_rejects_invalid_numbers(client, operator_headers, sample_product):
    product_id = sample_product["id"]
    rows = [
        ["周期开始 *", "周期结束 *", "曝光量", "点击量", "成交量", "花费金额", "成交金额",
         "点击率(%)", "转化率(%)", "ROI", "备注"],
        ["2026-09-01", "2026-09-07", 80000, 2400, 48, 1200.0, 4320.0, None, None, None, "有效行"],
        ["日期格式错误", "2026-09-14", 1000, 100, 1, 10.0, 20.0, None, None, None, "日期非法"],
        ["2026-10-01", "2026-10-07", 5000, 9000, 1, 10.0, 20.0, None, None, None, "点击大于曝光"],
        ["2026-10-08", "2026-09-01", 100, 10, 1, 1.0, 1.0, None, None, None, "结束早于开始"],
    ]
    response = client.post(
        f"/api/v1/products/{product_id}/performance-records/import?preview_only=true",
        files={"file": ("perf.xlsx", _xlsx(rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=operator_headers,
    )
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["valid_rows"] == 1
    assert data["error_rows"] == 3

    messages = " ".join(e for row in data["invalid_rows"] for e in row["errors"])
    assert "格式应为 YYYY-MM-DD" in messages
    assert "不能大于曝光量" in messages
    assert "早于周期开始" in messages


def test_import_requires_operator_role(client, viewer_headers, sample_store):
    rows = [PRODUCT_HEADER, ["测试店铺", "viewer不该能导入", "taobao", "箱包", 10.0, 5.0, "", "", "", "draft"]]
    response = client.post(
        "/api/v1/imports/products/preview",
        files={"file": ("p.xlsx", _xlsx(rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=viewer_headers,
    )
    assert response.status_code == 403
